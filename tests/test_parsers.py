import json

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from app.core.errors import HubError
from app.ingestion.chunking import chunk_blocks
from app.ingestion.parsers import parse, validate_file


def test_pdf_page_and_bbox_are_preserved(tmp_path):
    file = tmp_path / "source.pdf"
    pdf = pymupdf.open()
    pdf.new_page().insert_text((50, 70), "Evidence from page one.")
    pdf.save(file)
    pdf.close()
    blocks, parser = parse(file)
    chunks = chunk_blocks(blocks, "version-1")
    assert chunks[0]["page"] == 1 and len(chunks[0]["bbox"]) == 4
    assert "Evidence" in chunks[0]["text"]


def test_docx_and_xlsx_tables_are_not_flattened_into_unrelated_paragraphs(tmp_path):
    word = tmp_path / "table.docx"
    doc = Document()
    doc.add_paragraph("Table introduction")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Unit"
    table.cell(1, 0).text, table.cell(1, 1).text = "Item", "mg"
    doc.save(word)
    assert any(block["kind"] == "table" and "Item | mg" in block["text"] for block in parse(word)[0])
    excel = tmp_path / "table.xlsx"
    workbook = Workbook()
    workbook.active.append(["Name", "Unit"])
    workbook.active.append(["Item", "mg"])
    workbook.save(excel)
    assert parse(excel)[0][0]["kind"] == "table"


def test_scanned_pdf_requires_parser_instead_of_succeeding_with_empty_text(tmp_path):
    file = tmp_path / "scan.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(file)
    pdf.close()
    with pytest.raises(HubError, match="MinerU"):
        parse(file)


def test_heading_parent_and_html_script_removal(tmp_path):
    file = tmp_path / "source.html"
    file.write_text("<script>steal-token</script><h1>Evidence</h1><p>Body</p>", encoding="utf-8")
    assert "steal-token" not in parse(file)[0][0]["text"]
    file = tmp_path / "source.md"
    file.write_text("# Title\n\n## Section\n" + "内容。" * 600, encoding="utf-8")
    chunks = chunk_blocks(parse(file)[0], "v", size=100, overlap=20)
    assert len(chunks) > 5 and chunks[1]["section"] == "Title / Section"
    assert len(chunks[1]["parent_text"]) > len(chunks[1]["text"])
    with pytest.raises(HubError):
        validate_file("not-a.pdf", b"not a PDF", 1000)
