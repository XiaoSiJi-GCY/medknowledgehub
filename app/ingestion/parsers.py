"""Parse into source-located blocks. Never silently claim OCR for a textless PDF."""

import json
from pathlib import Path
from zipfile import ZipFile

import pymupdf as fitz
from bs4 import BeautifulSoup
from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from app.core.errors import HubError

ALLOWED = {".pdf", ".md", ".txt", ".html", ".docx", ".xlsx", ".json"}


def validate_file(filename, content, limit):
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise HubError("支持 PDF、Markdown、TXT、HTML、DOCX、XLSX 和结构化 JSON")
    if not content or len(content) > limit:
        raise HubError("文件为空或超过上传限制", 413)
    if suffix == ".pdf" and not content.startswith(b"%PDF-"):
        raise HubError("文件内容与 PDF 类型不符")
    if suffix in {".docx", ".xlsx"}:
        from io import BytesIO

        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                    raise HubError("Office 文件解压后过大", 413)
        except HubError:
            raise
        except Exception as exc:
            raise HubError("Office 文件格式不正确") from exc
    return suffix


def block(text, page=None, bbox=None, kind="text", section=""):
    return {"text": text.strip(), "page": page, "bbox": bbox, "kind": kind, "section": section}


def parse(path: Path, mineru=None):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with fitz.open(path) as document:
            blocks = [
                block(item[4], page.number + 1, list(item[:4]))
                for page in document
                for item in page.get_text("blocks")
                if len(item) > 4 and str(item[4]).strip()
            ]
            if not blocks:
                if mineru:
                    return mineru.parse(path), "mineru-v1"
                raise HubError("PDF 没有可提取文本。请配置 MinerU 服务解析扫描件。")
            return blocks, "pymupdf-" + fitz.VersionBind
    if suffix == ".docx":
        doc = WordDocument(path)
        blocks, section = [], ""
        for child in doc.element.body:
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, doc)
                if paragraph.style and paragraph.style.name.startswith("Heading"):
                    section = paragraph.text
                if paragraph.text.strip():
                    blocks.append(block(paragraph.text, section=section))
            elif child.tag.endswith("}tbl"):
                table = Table(child, doc)
                blocks.append(
                    block(
                        "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows),
                        kind="table",
                        section=section,
                    )
                )
        return blocks, "python-docx"
    if suffix == ".xlsx":
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            blocks = []
            for sheet in workbook:
                rows = [" | ".join("" if value is None else str(value) for value in row) for row in sheet.values]
                blocks.append(block("\n".join(rows), kind="table", section=sheet.title))
            return blocks, "openpyxl"
        finally:
            workbook.close()
    text = path.read_text(encoding="utf-8-sig")
    if suffix == ".html":
        soup = BeautifulSoup(text, "html.parser")
        for item in soup(["script", "style", "noscript"]):
            item.decompose()
        text = soup.get_text("\n", strip=True)
    if suffix == ".json":
        data = json.loads(text)
        records = data.get("records", []) if isinstance(data, dict) else data
        if not isinstance(records, list) or len(records) > 2000:
            raise HubError("结构化 JSON 需要 records 数组，最多 2000 条")
        blocks = []
        for record in records:
            if not isinstance(record, dict) or not record.get("name"):
                raise HubError("每条结构化记录需要 name")
            limits = {"category": 50, "name": 200, "department": 100, "value": 200, "unit": 100, "description": 4000}
            allowed = {key: str(record.get(key, "")) for key in limits}
            if any(len(allowed[key]) > limit for key, limit in limits.items()):
                raise HubError("结构化字段过长，请按记录格式限制拆分内容")
            item = block(
                "\n".join(f"{key}: {value}" for key, value in allowed.items()), kind="record", section=allowed["name"]
            )
            item["record"] = allowed
            blocks.append(item)
        return blocks, "structured-json-v1"
    return [block(text)], "text-v1"
